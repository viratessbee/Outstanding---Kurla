"""Convert the 'Customerwise Outstanding' Excel report into data.json for the website.
Usage:  python convert.py OS_24_OCT_2026.xlsx
"""
import sys, json, re
from datetime import datetime
from openpyxl import load_workbook

src = sys.argv[1] if len(sys.argv) > 1 else "OS.xlsx"
ws = load_workbook(src, data_only=True).active

def d(s):
    if not s: return None
    if isinstance(s, datetime): return s.strftime("%Y-%m-%d")
    try: return datetime.strptime(str(s).strip(), "%d-%b-%Y").strftime("%Y-%m-%d")
    except ValueError: return None

def num(x):
    try: return round(float(x), 2)
    except (TypeError, ValueError): return 0

report_date = None
for r in ws.iter_rows(min_row=1, max_row=6, values_only=True):
    m = re.search(r"To Date:\s*(\d{2}-\w{3}-\d{4})", str(r[0] or ""))
    if m: report_date = d(m.group(1))

parties, cur = {}, None
for r in ws.iter_rows(min_row=1, values_only=True):
    a = r[0]
    if a is None: continue                      # total rows have empty Book
    a = str(a).strip()
    if a == "Book" or a.startswith(("ESSBEE", "Customerwise", "From Date")):
        continue                                # title / repeated page headers
    if all(v is None for v in r[1:]):           # party header line
        cur = a; parties.setdefault(cur, []); continue
    if cur is None: continue
    inv_amt, os_amt, net = num(r[3]), num(r[5]), num(r[9])
    parties[cur].append({
        "book": a, "inv": str(r[1] or ""), "inv_date": d(r[2]),
        "inv_amt": inv_amt, "received": num(r[4]), "os": os_amt,
        "unalloc": num(r[8]), "net": net,
        "dispatch": d(r[11]), "cd": str(r[12] or "").strip() == "#", "sales": (str(r[13]).strip() if r[13] else "UNKNOWN"),
    })

# Salesman = most common on the party's rows (a party sits with one salesperson)
out = []
for p, rows in parties.items():
    if not rows: continue
    from collections import Counter
    sp = Counter(x["sales"] for x in rows).most_common(1)[0][0]
    out.append({"party": p, "sales": sp, "rows": rows})

json.dump({"report_date": report_date, "parties": out}, open("data.json", "w"),
          separators=(",", ":"))
print("Report date:", report_date, "| parties:", len(out),
      "| rows:", sum(len(p["rows"]) for p in out))
